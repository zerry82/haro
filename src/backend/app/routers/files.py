from __future__ import annotations

import os
import csv
from io import BytesIO, StringIO

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.dependencies import get_current_user, get_db
from app.models.project import Project
from app.models.user import User
from app.services.harness import ensure_harness_structure, is_clean_room_path, is_haro_internal_path
from app.services.workspace_index import META_EXCLUDES, rebuild_workspace_index, update_file_summary

router = APIRouter(prefix="/api/projects/{project_id}/files", tags=["files"])


class FileItem(BaseModel):
    name: str
    type: str
    size: int | None = None
    children_count: int | None = None


class FileListResponse(BaseModel):
    path: str
    items: list[FileItem]


class FileContentResponse(BaseModel):
    path: str
    content: str
    size: int
    language: str


class FileContentUpdate(BaseModel):
    content: str


class DirectoryCreateRequest(BaseModel):
    path: str


class FileMutationResponse(BaseModel):
    path: str
    type: str


class FileUploadResponse(BaseModel):
    path: str
    uploaded: list[str]


LANG_MAP = {
    ".py": "python", ".js": "javascript", ".ts": "typescript",
    ".html": "html", ".css": "css", ".json": "json",
    ".csv": "csv",
    ".md": "markdown", ".yaml": "yaml", ".yml": "yaml",
    ".txt": "plaintext", ".svg": "xml",
}

EDITABLE_EXTENSIONS = {
    ".html", ".md", ".csv", ".ts", ".js", ".json",
    ".txt", ".css", ".py", ".yaml", ".yml", ".svg",
}

EXCEL_EXTENSIONS = {".xlsx", ".xlsm", ".xls"}


async def _get_workspace(db: AsyncSession, user_id: str, project_id: str) -> str:
    result = await db.execute(
        select(Project).where(Project.id == project_id, Project.user_id == user_id)
    )
    project = result.scalar_one_or_none()
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    return project.workspace_path


def _assert_read_allowed(requested: str) -> None:
    if is_haro_internal_path(requested):
        raise HTTPException(status_code=403, detail="Haro internal metadata is not accessible")


def _assert_write_allowed(requested: str) -> None:
    _assert_read_allowed(requested)
    if is_clean_room_path(requested):
        raise HTTPException(status_code=403, detail="Clean Room is read-only")


def _validate_path(workspace: str, requested: str) -> str:
    workspace_real = os.path.realpath(workspace)
    full = os.path.realpath(os.path.join(workspace_real, requested.lstrip("/")))
    try:
        inside_workspace = os.path.commonpath([workspace_real, full]) == workspace_real
    except ValueError:
        inside_workspace = False
    if not inside_workspace:
        raise HTTPException(status_code=403, detail="Path traversal denied")
    return full


def _join_workspace_path(base_path: str, name: str) -> str:
    return f"/{name}" if base_path == "/" else f"{base_path.rstrip('/')}/{name}"


def _validate_file_name(name: str) -> None:
    if not name or name in {".", ".."} or "/" in name or "\\" in name:
        raise HTTPException(status_code=400, detail=f"Invalid file name: {name}")


def _sanitize_file_segment(value: str, fallback: str) -> str:
    cleaned = "".join(
        char if char not in '<>:"/\\|?*' and ord(char) >= 32 else "_"
        for char in value.strip()
    ).strip(" .")
    if not cleaned or cleaned in {".", ".."}:
        return fallback
    return cleaned


def _dedupe_name(name: str, used: set[str]) -> str:
    if name not in used:
        used.add(name)
        return name
    stem, ext = os.path.splitext(name)
    index = 2
    while True:
        candidate = f"{stem}_{index}{ext}"
        if candidate not in used:
            used.add(candidate)
            return candidate
        index += 1


def _stringify_excel_cell(value: object) -> str:
    return "" if value is None else str(value)


def _csv_bytes(rows: list[list[object]]) -> bytes:
    stream = StringIO(newline="")
    writer = csv.writer(stream, lineterminator="\n")
    writer.writerows([[_stringify_excel_cell(cell) for cell in row] for row in rows])
    return stream.getvalue().encode("utf-8")


def _excel_to_csv_outputs(filename: str, content: bytes) -> list[tuple[str, bytes]]:
    ext = os.path.splitext(filename)[1].lower()
    workbook_name = _sanitize_file_segment(os.path.splitext(filename)[0], "workbook")
    used_sheet_names: set[str] = set()
    outputs: list[tuple[str, bytes]] = []

    if ext in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise HTTPException(status_code=500, detail="Excel conversion dependency is missing") from exc

        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"Invalid Excel file: {filename}") from exc

        for sheet in workbook.worksheets:
            sheet_name = _sanitize_file_segment(sheet.title, "sheet")
            csv_name = _dedupe_name(f"{sheet_name}.csv", used_sheet_names)
            rows = [[cell for cell in row] for row in sheet.iter_rows(values_only=True)]
            outputs.append((_join_workspace_path(workbook_name, csv_name), _csv_bytes(rows)))
        workbook.close()
        return outputs

    if ext == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise HTTPException(status_code=500, detail="Excel conversion dependency is missing") from exc

        try:
            workbook = xlrd.open_workbook(file_contents=content)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"Invalid Excel file: {filename}") from exc

        for sheet in workbook.sheets():
            sheet_name = _sanitize_file_segment(sheet.name, "sheet")
            csv_name = _dedupe_name(f"{sheet_name}.csv", used_sheet_names)
            rows = [sheet.row_values(row_index) for row_index in range(sheet.nrows)]
            outputs.append((_join_workspace_path(workbook_name, csv_name), _csv_bytes(rows)))
        return outputs

    return []


def _is_excel_file(filename: str) -> bool:
    return os.path.splitext(filename)[1].lower() in EXCEL_EXTENSIONS


def _get_language(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    return LANG_MAP.get(ext, "plaintext")


@router.get("", response_model=FileListResponse)
async def list_files(
    project_id: str,
    path: str = Query("/"),
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    workspace = await _get_workspace(db, user.id, project_id)
    ensure_harness_structure(workspace, user.id, initialize_git=path == "/")
    _assert_read_allowed(path)
    full_path = _validate_path(workspace, path)
    if not os.path.isdir(full_path):
        raise HTTPException(status_code=404, detail="Directory not found")

    items: list[FileItem] = []
    for entry in os.scandir(full_path):
        if entry.name in META_EXCLUDES:
            continue
        if entry.is_dir():
            children = len([e for e in os.scandir(entry.path) if e.name not in META_EXCLUDES])
            items.append(FileItem(name=entry.name, type="directory", children_count=children))
        else:
            items.append(FileItem(name=entry.name, type="file", size=entry.stat().st_size))
    items.sort(key=lambda x: (x.type == "file", x.name))
    return FileListResponse(path=path, items=items)


@router.post("/directories", response_model=FileMutationResponse, status_code=201)
async def create_directory(
    project_id: str,
    body: DirectoryCreateRequest,
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    workspace = await _get_workspace(db, user.id, project_id)
    ensure_harness_structure(workspace, user.id, initialize_git=False)
    _assert_write_allowed(body.path)
    full_path = _validate_path(workspace, body.path)
    if os.path.exists(full_path):
        raise HTTPException(status_code=409, detail="Path already exists")

    parent = os.path.dirname(full_path)
    if not os.path.isdir(parent):
        raise HTTPException(status_code=404, detail="Parent directory not found")

    os.mkdir(full_path)
    await rebuild_workspace_index(workspace)
    return FileMutationResponse(path=body.path, type="directory")


@router.post("/upload", response_model=FileUploadResponse)
async def upload_files(
    project_id: str,
    path: str = Query("/"),
    overwrite: bool = Query(False),
    files: list[UploadFile] = File(...),
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    workspace = await _get_workspace(db, user.id, project_id)
    ensure_harness_structure(workspace, user.id, initialize_git=False)
    _assert_write_allowed(path)
    target_dir = _validate_path(workspace, path)
    if not os.path.isdir(target_dir):
        raise HTTPException(status_code=404, detail="Directory not found")

    seen: set[str] = set()
    upload_items: list[tuple[str, bytes]] = []
    for upload in files:
        filename = upload.filename or ""
        _validate_file_name(filename)
        if filename in seen:
            raise HTTPException(status_code=400, detail=f"Duplicate upload name: {filename}")
        seen.add(filename)
        content = await upload.read()
        if _is_excel_file(filename):
            for relative_name, csv_content in _excel_to_csv_outputs(filename, content):
                upload_items.append((relative_name, csv_content))
        else:
            upload_items.append((filename, content))

    generated_paths: set[str] = set()
    existed_before: dict[str, bool] = {}
    for relative_name, _content in upload_items:
        relative_path = _join_workspace_path(path, relative_name)
        target_path = _validate_path(workspace, relative_path)
        if target_path in generated_paths:
            raise HTTPException(status_code=400, detail=f"Duplicate upload output: {relative_name}")
        generated_paths.add(target_path)

        parent = os.path.dirname(target_path)
        if os.path.exists(parent) and not os.path.isdir(parent):
            raise HTTPException(status_code=409, detail=f"File already exists: {os.path.basename(parent)}")
        if os.path.isdir(target_path):
            raise HTTPException(status_code=409, detail=f"Directory already exists: {relative_name}")
        existed_before[relative_path] = os.path.exists(target_path)
        if os.path.exists(target_path) and not overwrite:
            raise HTTPException(status_code=409, detail=f"File already exists: {relative_name}")

    uploaded: list[str] = []
    for relative_name, content in upload_items:
        relative_path = _join_workspace_path(path, relative_name)
        target_path = _validate_path(workspace, relative_path)
        os.makedirs(os.path.dirname(target_path), exist_ok=True)
        with open(target_path, "wb") as f:
            f.write(content)
        uploaded.append(relative_path)
        await update_file_summary(workspace, relative_path, "modified" if existed_before.get(relative_path) else "created")

    return FileUploadResponse(path=path, uploaded=uploaded)


@router.get("/content", response_model=FileContentResponse)
async def get_file_content(
    project_id: str,
    path: str = Query(...),
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    workspace = await _get_workspace(db, user.id, project_id)
    ensure_harness_structure(workspace, user.id, initialize_git=False)
    _assert_read_allowed(path)
    full_path = _validate_path(workspace, path)
    if not os.path.isfile(full_path):
        raise HTTPException(status_code=404, detail="File not found")

    with open(full_path, "r", encoding="utf-8", errors="replace") as f:
        content = f.read()

    language = _get_language(path)
    return FileContentResponse(path=path, content=content, size=len(content), language=language)


@router.put("/content", response_model=FileContentResponse)
async def update_file_content(
    project_id: str,
    body: FileContentUpdate,
    path: str = Query(...),
    user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    workspace = await _get_workspace(db, user.id, project_id)
    ensure_harness_structure(workspace, user.id, initialize_git=False)
    _assert_write_allowed(path)
    full_path = _validate_path(workspace, path)
    if not os.path.isfile(full_path):
        raise HTTPException(status_code=404, detail="File not found")

    ext = os.path.splitext(path)[1].lower()
    if ext not in EDITABLE_EXTENSIONS:
        raise HTTPException(status_code=415, detail="File type is not editable")

    with open(full_path, "w", encoding="utf-8") as f:
        f.write(body.content)

    await update_file_summary(workspace, path, "modified")

    language = _get_language(path)
    content = body.content
    return FileContentResponse(path=path, content=content, size=len(content), language=language)
