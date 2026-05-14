from __future__ import annotations

import asyncio
import csv
import json
import os
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any, Callable

from app.services.file_conversion import dedupe_name, sanitize_file_segment
from app.services.mail_management import attachment_signals, decide_management
from app.services.workspace_file_db import atomic_write_bytes, sync_workspace_path
from app.services.workspace_file_helpers import full_path
from app.services.workspace_index import update_file_summary
from app.services.llm import get_client

TEXT_EXTENSIONS = {"txt", "md"}
CSV_EXTENSIONS = {"csv"}
EXCEL_EXTENSIONS = {"xlsx", "xls"}
PDF_EXTENSIONS = {"pdf"}
IMAGE_EXTENSIONS = {"png", "jpeg", "jpg"}
SUMMARIZABLE_EXTENSIONS = TEXT_EXTENSIONS | CSV_EXTENSIONS | EXCEL_EXTENSIONS | PDF_EXTENSIONS | IMAGE_EXTENSIONS

MAX_TEXT_PREVIEW = 2000
MAX_KEY_POINTS = 5


@dataclass(frozen=True)
class AttachmentProcessingStats:
    downloaded_count: int = 0
    summarized_count: int = 0
    failed_count: int = 0


class MailAttachmentProcessor:
    def __init__(
        self,
        *,
        client_factory: Callable[[], Any] = get_client,
        image_model: str = "gemini-3-flash-preview",
    ) -> None:
        self.client_factory = client_factory
        self.image_model = image_model

    async def process_threads(
        self,
        *,
        workspace_path: str,
        user_id: str,
        token_ref: str | None,
        account_email: str,
        threads: list[dict[str, Any]],
        gmail_fetcher: Any,
        management_policy: dict[str, Any] | None,
    ) -> AttachmentProcessingStats:
        downloaded_count = 0
        summarized_count = 0
        failed_count = 0

        for thread in threads:
            attachments = thread.get("attachments") if isinstance(thread.get("attachments"), list) else []
            if not attachments:
                continue
            if _skip_for_management(thread, attachments, management_policy):
                thread["attachments"] = [_skipped_attachment(attachment) for attachment in attachments if isinstance(attachment, dict)]
                continue

            processed: list[dict[str, Any]] = []
            for attachment in attachments:
                if not isinstance(attachment, dict):
                    continue
                row = dict(attachment)
                try:
                    content = await self._ensure_downloaded(
                        workspace_path=workspace_path,
                        user_id=user_id,
                        thread=thread,
                        attachment=row,
                        token_ref=token_ref,
                        account_email=account_email,
                        gmail_fetcher=gmail_fetcher,
                    )
                    if row.get("download_status") == "downloaded":
                        downloaded_count += 1
                    row = await self._extract(row, content)
                    if row.get("summary"):
                        summarized_count += 1
                    if row.get("extract_status") == "failed":
                        failed_count += 1
                except Exception as exc:
                    row["extract_status"] = "failed"
                    row["error_reason"] = str(exc) or "attachment processing failed"
                    failed_count += 1
                processed.append(row)
            thread["attachments"] = processed

        return AttachmentProcessingStats(downloaded_count, summarized_count, failed_count)

    async def _ensure_downloaded(
        self,
        *,
        workspace_path: str,
        user_id: str,
        thread: dict[str, Any],
        attachment: dict[str, Any],
        token_ref: str | None,
        account_email: str,
        gmail_fetcher: Any,
    ) -> bytes:
        inbox_path = str(attachment.get("inbox_path") or "").strip()
        if inbox_path:
            try:
                full = full_path(workspace_path, inbox_path)
                if os.path.exists(full):
                    attachment["download_status"] = "cached"
                    return await asyncio.to_thread(Path(full).read_bytes)
            except Exception:
                pass

        message_id = str(attachment.get("_gmail_message_id") or "").strip()
        attachment_id = str(attachment.get("_gmail_attachment_id") or "").strip()
        if not token_ref or not message_id or not attachment_id:
            attachment["download_status"] = "missing_ref"
            return b""

        content = await gmail_fetcher.download_attachment(
            token_ref=token_ref,
            message_id=message_id,
            attachment_id=attachment_id,
        )
        target_path = _attachment_inbox_path(user_id, thread, attachment)
        target_path = _dedupe_workspace_file_path(workspace_path, target_path)
        full = full_path(workspace_path, target_path)
        await asyncio.to_thread(atomic_write_bytes, full, content)
        sync_workspace_path(workspace_path, target_path, source_kind="mail_attachment")
        await update_file_summary(workspace_path, target_path, "created")
        attachment["inbox_path"] = target_path
        attachment["download_status"] = "downloaded"
        attachment["downloaded_at"] = _compact(str(thread.get("received_at") or ""), 50)
        return content

    async def _extract(self, attachment: dict[str, Any], content: bytes) -> dict[str, Any]:
        extension = _extension(attachment)
        if extension not in SUMMARIZABLE_EXTENSIONS:
            attachment.setdefault("extract_status", "metadata_only")
            attachment.setdefault("summary", "")
            attachment.setdefault("content_profile", {"kind": "metadata_only"})
            return attachment
        if not content:
            attachment["extract_status"] = "failed"
            attachment["error_reason"] = attachment.get("error_reason") or "downloaded attachment content is empty"
            return attachment

        if extension in TEXT_EXTENSIONS:
            profile = _text_profile(content)
        elif extension in CSV_EXTENSIONS:
            profile = _csv_profile(content)
        elif extension in EXCEL_EXTENSIONS:
            profile = _excel_profile(extension, content)
        elif extension in PDF_EXTENSIONS:
            profile = _pdf_profile(content)
        else:
            profile = await asyncio.to_thread(self._image_profile, attachment, content)

        attachment["content_profile"] = profile
        if profile.get("error_reason"):
            attachment["extract_status"] = "failed"
            attachment["error_reason"] = str(profile.get("error_reason") or "attachment extraction failed")
            return attachment
        attachment["summary"] = _summary_from_profile(attachment, profile)
        attachment["key_points"] = _key_points_from_profile(profile)
        if profile.get("summary_status") == "unsupported_scanned_pdf":
            attachment["extract_status"] = "unsupported_scanned_pdf"
        else:
            attachment["extract_status"] = "summarized" if attachment["summary"] else "metadata_only"
        return attachment

    def _image_profile(self, attachment: dict[str, Any], content: bytes) -> dict[str, Any]:
        mime_type = str(attachment.get("mime_type") or _image_mime(_extension(attachment)) or "image/png")
        prompt = (
            "이 Gmail 이미지 첨부파일을 한국어 업무 맥락으로 요약하세요. "
            "보이는 텍스트, 표/차트/스크린샷 여부, 확인해야 할 사항을 간결히 JSON으로 반환하세요. "
            "Schema: {\"image_description\":\"string\",\"visible_text\":\"string\",\"key_points\":[\"string\"]}"
        )
        try:
            from google.genai import types

            response = self.client_factory().models.generate_content(
                model=self.image_model,
                contents=[
                    prompt,
                    types.Part.from_bytes(data=content, mime_type=mime_type),
                ],
                config={"temperature": 0.1, "response_mime_type": "application/json"},
            )
            text = str(getattr(response, "text", "") or "")
            parsed = _parse_json_object(text)
        except Exception as exc:
            parsed = {"image_description": "", "visible_text": "", "key_points": [], "error_reason": str(exc)}
        return {
            "kind": "image",
            "image_description": str(parsed.get("image_description") or ""),
            "visible_text": str(parsed.get("visible_text") or ""),
            "key_points": [str(item) for item in parsed.get("key_points") or []][:MAX_KEY_POINTS],
            "error_reason": str(parsed.get("error_reason") or ""),
        }


def _skip_for_management(thread: dict[str, Any], attachments: list[dict[str, Any]], policy: dict[str, Any] | None) -> bool:
    metadata = dict(thread.get("metadata") if isinstance(thread.get("metadata"), dict) else {})
    metadata["attachment_signals"] = attachment_signals([attachment for attachment in attachments if isinstance(attachment, dict)])
    decision = decide_management(thread, policy, metadata)
    return decision.decision == "excluded" and decision.source in {"blacklist", "manual_override"}


def _skipped_attachment(attachment: dict[str, Any]) -> dict[str, Any]:
    row = dict(attachment)
    row["download_status"] = "skipped_excluded"
    row.setdefault("extract_status", "skipped")
    return row


def _attachment_inbox_path(user_id: str, thread: dict[str, Any], attachment: dict[str, Any]) -> str:
    date = str(thread.get("received_at") or "")[:10] or "unknown-date"
    subject = sanitize_file_segment(str(thread.get("subject") or thread.get("source_ref") or "mail"), "mail")
    filename = sanitize_file_segment(str(attachment.get("title") or "attachment"), "attachment")
    return f"/playground/users/{user_id}/00_inbox/mail/gmail/{date}/{subject}/{filename}"


def _dedupe_workspace_file_path(workspace_path: str, requested_path: str) -> str:
    directory, filename = requested_path.rsplit("/", 1)
    stem, ext = os.path.splitext(filename)
    used = set()
    candidate = filename
    index = 2
    while os.path.exists(full_path(workspace_path, f"{directory}/{candidate}")):
        used.add(candidate)
        candidate = dedupe_name(f"{stem}_{index}{ext}", used)
        index += 1
    return f"{directory}/{candidate}"


def _text_profile(content: bytes) -> dict[str, Any]:
    text = content.decode("utf-8-sig", errors="replace")
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return {
        "kind": "text",
        "line_count": len(lines),
        "text_preview": _compact("\n".join(lines), MAX_TEXT_PREVIEW),
    }


def _csv_profile(content: bytes) -> dict[str, Any]:
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.reader(StringIO(text))
    rows = [row for _, row in zip(range(101), reader)]
    header = rows[0] if rows else []
    data_rows = rows[1:]
    return {
        "kind": "csv",
        "columns": [str(item) for item in header],
        "row_count_sampled": len(data_rows),
        "sample_rows": data_rows[:5],
        "text_preview": _compact(text, MAX_TEXT_PREVIEW),
    }


def _excel_profile(extension: str, content: bytes) -> dict[str, Any]:
    if extension == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            sheets = []
            for sheet in workbook.worksheets[:10]:
                rows = []
                for index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if index >= 6:
                        break
                    rows.append([_cell_text(cell) for cell in row])
                sheets.append(_sheet_profile(sheet.title, rows, sheet.max_row, sheet.max_column))
            return {"kind": "excel", "sheets": sheets, "text_preview": _compact(json.dumps(sheets, ensure_ascii=False), MAX_TEXT_PREVIEW)}
        finally:
            workbook.close()

    import xlrd

    workbook = xlrd.open_workbook(file_contents=content)
    sheets = []
    for sheet in workbook.sheets()[:10]:
        rows = [[_cell_text(cell) for cell in sheet.row_values(row_index)] for row_index in range(min(sheet.nrows, 6))]
        sheets.append(_sheet_profile(sheet.name, rows, sheet.nrows, sheet.ncols))
    return {"kind": "excel", "sheets": sheets, "text_preview": _compact(json.dumps(sheets, ensure_ascii=False), MAX_TEXT_PREVIEW)}


def _sheet_profile(name: str, rows: list[list[str]], row_count: int, column_count: int) -> dict[str, Any]:
    columns = rows[0] if rows else []
    return {
        "name": name,
        "row_count": row_count,
        "column_count": column_count,
        "columns": columns,
        "sample_rows": rows[1:6],
    }


def _pdf_profile(content: bytes) -> dict[str, Any]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("pypdf dependency is missing") from exc
    reader = PdfReader(BytesIO(content))
    text = "\n".join(page.extract_text() or "" for page in reader.pages[:20]).strip()
    if len(text) < 40:
        return {"kind": "pdf", "page_count": len(reader.pages), "text_preview": "", "summary_status": "unsupported_scanned_pdf"}
    return {"kind": "pdf", "page_count": len(reader.pages), "text_preview": _compact(text, MAX_TEXT_PREVIEW)}


def _summary_from_profile(attachment: dict[str, Any], profile: dict[str, Any]) -> str:
    title = str(attachment.get("title") or "첨부파일")
    kind = str(profile.get("kind") or "")
    if kind == "csv":
        columns = ", ".join(str(item) for item in profile.get("columns") or [] if str(item))
        return f"{title}: CSV 파일입니다. 컬럼은 {columns or '확인되지 않음'}이며 샘플 기준 {profile.get('row_count_sampled', 0)}개 데이터 행을 확인했습니다."
    if kind == "excel":
        sheets = profile.get("sheets") if isinstance(profile.get("sheets"), list) else []
        names = ", ".join(str(sheet.get("name")) for sheet in sheets if isinstance(sheet, dict))
        return f"{title}: Excel 파일입니다. 시트 {len(sheets)}개({names or '이름 없음'})의 표 구조와 샘플 데이터를 확인했습니다."
    if kind == "pdf":
        if profile.get("summary_status") == "unsupported_scanned_pdf":
            return f"{title}: PDF 파일이지만 추출 가능한 텍스트가 거의 없어 스캔본으로 보입니다."
        return f"{title}: PDF 파일입니다. {profile.get('page_count', 0)}페이지에서 텍스트를 추출했습니다. {_compact(str(profile.get('text_preview') or ''), 240)}"
    if kind == "image":
        description = str(profile.get("image_description") or "")
        visible_text = str(profile.get("visible_text") or "")
        return f"{title}: 이미지 파일입니다. {description or visible_text or '이미지 내용을 확인했습니다.'}"
    preview = str(profile.get("text_preview") or "")
    return f"{title}: 텍스트 파일입니다. {_compact(preview, 240)}" if preview else ""


def _key_points_from_profile(profile: dict[str, Any]) -> list[str]:
    if isinstance(profile.get("key_points"), list):
        return [str(item) for item in profile["key_points"][:MAX_KEY_POINTS] if str(item).strip()]
    rows: list[str] = []
    if profile.get("columns"):
        rows.append("컬럼: " + ", ".join(str(item) for item in profile.get("columns") or [] if str(item)))
    if profile.get("sheets"):
        rows.append(f"시트 수: {len(profile.get('sheets') or [])}")
    if profile.get("page_count"):
        rows.append(f"PDF 페이지 수: {profile.get('page_count')}")
    preview = str(profile.get("text_preview") or "").strip()
    if preview:
        rows.append(_compact(preview, 180))
    return rows[:MAX_KEY_POINTS]


def _extension(attachment: dict[str, Any]) -> str:
    title = str(attachment.get("title") or "")
    return str(attachment.get("extension") or os.path.splitext(title)[1].lstrip(".")).casefold()


def _image_mime(extension: str) -> str:
    if extension in {"jpg", "jpeg"}:
        return "image/jpeg"
    if extension == "png":
        return "image/png"
    return "image/png"


def _parse_json_object(text: str) -> dict[str, Any]:
    cleaned = str(text or "").strip()
    if cleaned.startswith("```"):
        cleaned = cleaned.strip("`").strip()
        if cleaned.startswith("json"):
            cleaned = cleaned[4:].strip()
    start = cleaned.find("{")
    end = cleaned.rfind("}")
    if start < 0 or end < start:
        return {}
    try:
        parsed = json.loads(cleaned[start:end + 1])
    except json.JSONDecodeError:
        return {}
    return parsed if isinstance(parsed, dict) else {}


def _cell_text(value: Any) -> str:
    return "" if value is None else str(value)


def _compact(value: str, limit: int) -> str:
    compact = " ".join(str(value or "").split())
    if len(compact) <= limit:
        return compact
    return f"{compact[:limit - 1].rstrip()}..."
