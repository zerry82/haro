from __future__ import annotations

import csv
import os
from io import BytesIO, StringIO

from fastapi import HTTPException

from app.services.file_path_policy import join_workspace_path


LANG_MAP = {
    ".py": "python",
    ".js": "javascript",
    ".ts": "typescript",
    ".html": "html",
    ".css": "css",
    ".json": "json",
    ".csv": "csv",
    ".md": "markdown",
    ".yaml": "yaml",
    ".yml": "yaml",
    ".txt": "plaintext",
    ".svg": "xml",
}

EDITABLE_EXTENSIONS = {
    ".html",
    ".md",
    ".csv",
    ".ts",
    ".js",
    ".json",
    ".txt",
    ".css",
    ".py",
    ".yaml",
    ".yml",
    ".svg",
}

EXCEL_EXTENSIONS = {".xlsx", ".xlsm", ".xls"}


def sanitize_file_segment(value: str, fallback: str) -> str:
    cleaned = "".join(
        char if char not in '<>:"/\\|?*' and ord(char) >= 32 else "_"
        for char in value.strip()
    ).strip(" .")
    if not cleaned or cleaned in {".", ".."}:
        return fallback
    return cleaned


def dedupe_name(name: str, used: set[str]) -> str:
    if name not in used:
        used.add(name)
        return name
    stem, ext = os.path.splitext(name)
    index = 2
    while True:
        candidate = f"{stem}_{index}{ext}"
        if candidate not in used:
            used.add(candidate)
            return candidate
        index += 1


def _stringify_excel_cell(value: object) -> str:
    return "" if value is None else str(value)


def _csv_bytes(rows: list[list[object]]) -> bytes:
    stream = StringIO(newline="")
    writer = csv.writer(stream, lineterminator="\n")
    writer.writerows([[_stringify_excel_cell(cell) for cell in row] for row in rows])
    return stream.getvalue().encode("utf-8")


def excel_to_csv_outputs(filename: str, content: bytes) -> list[tuple[str, bytes]]:
    ext = os.path.splitext(filename)[1].lower()
    workbook_name = sanitize_file_segment(os.path.splitext(filename)[0], "workbook")
    used_sheet_names: set[str] = set()
    outputs: list[tuple[str, bytes]] = []

    if ext in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise HTTPException(status_code=500, detail="Excel conversion dependency is missing") from exc

        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"Invalid Excel file: {filename}") from exc

        for sheet in workbook.worksheets:
            sheet_name = sanitize_file_segment(sheet.title, "sheet")
            csv_name = dedupe_name(f"{sheet_name}.csv", used_sheet_names)
            rows = [[cell for cell in row] for row in sheet.iter_rows(values_only=True)]
            outputs.append((join_workspace_path(workbook_name, csv_name), _csv_bytes(rows)))
        workbook.close()
        return outputs

    if ext == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise HTTPException(status_code=500, detail="Excel conversion dependency is missing") from exc

        try:
            workbook = xlrd.open_workbook(file_contents=content)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"Invalid Excel file: {filename}") from exc

        for sheet in workbook.sheets():
            sheet_name = sanitize_file_segment(sheet.name, "sheet")
            csv_name = dedupe_name(f"{sheet_name}.csv", used_sheet_names)
            rows = [sheet.row_values(row_index) for row_index in range(sheet.nrows)]
            outputs.append((join_workspace_path(workbook_name, csv_name), _csv_bytes(rows)))
        return outputs

    return []


def is_excel_file(filename: str) -> bool:
    return os.path.splitext(filename)[1].lower() in EXCEL_EXTENSIONS


def get_language(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    return LANG_MAP.get(ext, "plaintext")
